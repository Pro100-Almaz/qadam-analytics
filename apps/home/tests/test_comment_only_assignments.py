"""
Spec 0005, AC-26 … AC-31 — comment-only assignments: `max_grade` is null, so
their grades carry comments and never a mark. Each test is named after the AC
it pins down.
"""

import datetime
import io

import pytest
from django.core.exceptions import ValidationError
from django.urls import reverse
from openpyxl import load_workbook

from apps.home.models import HomeroomTeacherAssignment, SubjectGrade
from apps.lesson.homework_sync import mirror_of
from core.factories import (
    ClassGroupFactory, EnrollmentFactory, HomeworkFactory,
    SubjectAssignmentFactory, SubjectGradeFactory, SubjectOfferingFactory,
    TeacherFactory, TeachingAssignmentFactory,
)

ASSIGNMENTS_URL = reverse('home-api:subject-assignment-list-create')

D = datetime.date


def detail_url(assignment):
    return reverse('home-api:subject-assignment-detail', args=[assignment.pk])


def grades_url(assignment):
    return reverse(
        'home-api:subject-assignment-grade-list-create', args=[assignment.pk],
    )


def grade_url(grade):
    return reverse('home-api:subject-grade-detail', args=[grade.pk])


@pytest.fixture
def year(academic_year):
    academic_year.q1_start, academic_year.q1_end = D(2026, 9, 1), D(2026, 10, 31)
    academic_year.save()
    return academic_year


@pytest.fixture
def world(year):
    class_group = ClassGroupFactory(academic_year=year)
    offering = SubjectOfferingFactory(class_group=class_group)
    teaching_assignment = TeachingAssignmentFactory(offering=offering)
    student = EnrollmentFactory(class_group=class_group).student
    return {
        'year': year,
        'class_group': class_group,
        'offering': offering,
        'teaching_assignment': teaching_assignment,
        'teacher': teaching_assignment.teacher,
        'student': student,
    }


@pytest.fixture
def teacher_client(world, authenticated_client):
    return authenticated_client(world['teacher'].user)


def comment_only(world, **kwargs):
    return SubjectAssignmentFactory(
        offering=world['offering'], max_grade=None, category='lesson',
        date=D(2026, 9, 15), **kwargs,
    )


def create(client, offering, **body):
    payload = {
        'offering': offering.id, 'title': 'Essay feedback',
        'date': '2026-09-15', **body,
    }
    return client.post(ASSIGNMENTS_URL, payload, format='json')


# ── AC-26 ──

def test_ac26_create_accepts_null_max_grade(world, teacher_client):
    response = create(teacher_client, world['offering'], max_grade=None)

    assert response.status_code == 201, response.data
    assert response.data['max_grade'] is None


def test_ac26_create_without_max_grade_is_400(world, teacher_client):
    response = create(teacher_client, world['offering'])

    assert response.status_code == 400
    assert 'max_grade' in response.data


# ── AC-27 ──

def test_ac27_homework_create_with_null_max_grade_is_400(world, teacher_client):
    response = create(
        teacher_client, world['offering'], max_grade=None, category='homework',
    )

    assert response.status_code == 400
    assert 'max_grade' in response.data


def test_ac27_homework_patch_to_null_max_grade_is_400(world, teacher_client):
    homework = HomeworkFactory(teaching_assignment=world['teaching_assignment'])
    assignment = mirror_of(homework.pk)

    response = teacher_client.patch(
        detail_url(assignment), {'max_grade': None}, format='json',
    )

    assert response.status_code == 400
    assert 'max_grade' in response.data
    homework.refresh_from_db()
    assert homework.max_grade == 10


def test_ac27_model_refuses_comment_only_homework_mirror(world):
    homework = HomeworkFactory(teaching_assignment=world['teaching_assignment'])
    assignment = mirror_of(homework.pk)
    assignment.max_grade = None

    with pytest.raises(ValidationError):
        assignment.save()


# ── AC-28 ──

def test_ac28_grade_on_comment_only_assignment_is_400(world, teacher_client):
    assignment = comment_only(world)

    response = teacher_client.post(
        grades_url(assignment),
        {'student': world['student'].id, 'grade': 5, 'comments': 'Nice'},
        format='json',
    )

    assert response.status_code == 400
    assert 'grade' in response.data
    assert not SubjectGrade.objects.filter(assignment=assignment).exists()


def test_ac28_comment_without_grade_is_accepted(world, teacher_client):
    assignment = comment_only(world)

    response = teacher_client.post(
        grades_url(assignment),
        {'student': world['student'].id, 'grade': None,
         'comments': 'Clear argument, cite your sources.'},
        format='json',
    )

    assert response.status_code == 201, response.data
    assert response.data['grade'] is None
    assert response.data['comments'] == 'Clear argument, cite your sources.'
    assert response.data['assignment']['max_grade'] is None


def test_ac28_patching_a_grade_onto_comment_only_is_400(world, teacher_client):
    assignment = comment_only(world)
    grade = SubjectGradeFactory(
        assignment=assignment, student=world['student'], comments='Good',
    )

    response = teacher_client.patch(grade_url(grade), {'grade': 3}, format='json')

    assert response.status_code == 400
    grade.refresh_from_db()
    assert grade.grade is None


def test_ac28_model_refuses_a_grade_on_comment_only(world):
    assignment = comment_only(world)

    with pytest.raises(ValidationError):
        SubjectGrade.objects.create(
            assignment=assignment, student=world['student'], grade=4,
        )


# ── AC-29 ──

def test_ac29_patch_to_null_with_marks_recorded_is_400(world, teacher_client):
    assignment = SubjectAssignmentFactory(
        offering=world['offering'], max_grade=10, category='lesson',
    )
    SubjectGradeFactory(assignment=assignment, student=world['student'], grade=7)

    response = teacher_client.patch(
        detail_url(assignment), {'max_grade': None}, format='json',
    )

    assert response.status_code == 400
    assert 'max_grade' in response.data
    assignment.refresh_from_db()
    assert assignment.max_grade == 10


def test_ac29_patch_to_null_without_marks_succeeds(world, teacher_client):
    assignment = SubjectAssignmentFactory(
        offering=world['offering'], max_grade=10, category='lesson',
    )
    # A comment with no mark does not block it.
    SubjectGradeFactory(
        assignment=assignment, student=world['student'], comments='Seen',
    )

    response = teacher_client.patch(
        detail_url(assignment), {'max_grade': None}, format='json',
    )

    assert response.status_code == 200, response.data
    assert response.data['max_grade'] is None


def test_ac29_patch_number_makes_comment_only_gradable(world, teacher_client):
    assignment = comment_only(world)

    response = teacher_client.patch(
        detail_url(assignment), {'max_grade': 20}, format='json',
    )

    assert response.status_code == 200, response.data
    assert response.data['max_grade'] == 20
    response = teacher_client.post(
        grades_url(assignment),
        {'student': world['student'].id, 'grade': 18}, format='json',
    )
    assert response.status_code == 201, response.data


# ── AC-30 ──

def test_ac30_analytics_leave_comment_only_work_out(
    world, authenticated_client,
):
    world['year'].is_active = True
    world['year'].save()
    graded = SubjectAssignmentFactory(
        offering=world['offering'], max_grade=10, category='lesson',
        date=D(2026, 9, 10),
    )
    SubjectGradeFactory(assignment=graded, student=world['student'], grade=8)
    feedback = comment_only(world)
    SubjectGradeFactory(
        assignment=feedback, student=world['student'], comments='Well argued',
    )
    client = authenticated_client(world['teacher'].user)
    student = world['student']

    trajectory = client.get(
        reverse(
            'lesson-api:analytics-assignment-trajectory',
            args=[student.id, world['offering'].id],
        ),
        {'missing': 'zero'},
    )
    assert trajectory.status_code == 200, trajectory.data
    assert [p['id'] for p in trajectory.data['points']] == [graded.id]
    assert trajectory.data['summary']['student_mean'] == 80.0
    assert trajectory.data['summary']['coverage']['possible_count'] == 1

    summary = client.get(
        reverse('lesson-api:analytics-assignment-summary', args=[student.id]),
        {'missing': 'zero', 'academic_year': world['year'].id},
    )
    assert summary.status_code == 200, summary.data
    axis = summary.data['axes'][0]
    assert axis['assignment_count'] == 1
    assert axis['value'] == 80.0


# ── AC-31 ──

def test_ac31_grade_sheet_marks_comment_only_columns(
    world, authenticated_client,
):
    feedback = comment_only(world, title='Essay')
    SubjectGradeFactory(
        assignment=feedback, student=world['student'], comments='Well argued',
    )
    homeroom = TeacherFactory()
    HomeroomTeacherAssignment.objects.create(
        teacher=homeroom, class_group=world['class_group'],
    )

    response = authenticated_client(homeroom.user).get(
        reverse(
            'student-report-api:class-group-grade-sheet',
            kwargs={'class_group_id': world['class_group'].pk},
        ),
        {'quarter': 1},
    )

    assert response.status_code == 200
    ws = load_workbook(io.BytesIO(response.content)).worksheets[0]
    header = ws.cell(row=4, column=2)
    assert header.value.startswith('Essay')
    assert header.comment.text.endswith('Comment only, no grade')
    cell = ws.cell(row=5, column=2)
    assert cell.value == '-'
    assert cell.comment.text.endswith('Well argued')
